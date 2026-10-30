
import csv
import io
import os

from flask import Blueprint, request, send_file, current_app
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from itsdangerous import (
    URLSafeTimedSerializer,
    BadSignature,
    SignatureExpired,
)
from sqlalchemy import or_

from extensions import db

from models.user import User
from models.hire_requirement import HireRequirement
from models.hire_file import HireFile

from models.deliver_proposal import DeliverProposal
from models.deliver_team_member import DeliverTeamMember
from models.deliver_file import DeliverFile

from services.auth import admin_required


admin_bp = Blueprint("admin", __name__)


# =========================================================
# CONSTANTS
# =========================================================

ALLOWED_STATUSES = [
    "Pending",
    "Approved",
    "In Progress",
    "Completed",
    "Rejected",
]

BASE_DIR = os.path.dirname(
    os.path.dirname(
        os.path.abspath(__file__)
    )
)

FILE_LINK_MAX_AGE = 7 * 24 * 60 * 60


# =========================================================
# GENERAL HELPERS
# =========================================================

def safe_text(value):
    """
    Safely convert any value to string.
    None becomes empty string.
    """
    if value is None:
        return ""

    return str(value)


def format_datetime(value):
    """
    Format database datetime values consistently.
    """
    if not value:
        return ""

    try:
        return value.strftime(
            "%Y-%m-%d %H:%M:%S"
        )
    except Exception:
        return str(value)


def parse_page():
    """
    Read and sanitize pagination parameters.
    """
    page = request.args.get(
        "page",
        1,
        type=int
    )

    per_page = request.args.get(
        "per_page",
        10,
        type=int
    )

    page = max(page, 1)

    per_page = max(
        1,
        min(per_page, 100)
    )

    return page, per_page


def normalize_record_type(value):
    """
    Allow only:
    all
    hire
    deliver
    """
    value = safe_text(value).strip().lower()

    if value in {
        "all",
        "hire",
        "deliver"
    }:
        return value

    return "all"


def get_user_info(user):
    """
    Return safe user information for admin dashboard.
    """
    if not user:
        return {
            "id": None,
            "name": "",
            "email": "",
            "phone": "",
            "role": "",
        }

    return {
        "id": user.id,
        "name": safe_text(user.name),
        "email": safe_text(user.email),
        "phone": safe_text(user.phone),
        "role": safe_text(user.role),
    }


# =========================================================
# FILE HELPERS
# =========================================================

def get_file_path(file_record):
    """
    Resolve physical file path safely.

    Prevents files from being accessed outside
    the backend directory.
    """
    if not file_record:
        return None

    raw_path = getattr(
        file_record,
        "file_path",
        None
    )

    if not raw_path:
        return None

    if os.path.isabs(raw_path):
        path = os.path.abspath(raw_path)
    else:
        path = os.path.abspath(
            os.path.join(
                BASE_DIR,
                raw_path
            )
        )

    try:
        if os.path.commonpath(
            [path, BASE_DIR]
        ) != BASE_DIR:
            return None
    except ValueError:
        return None

    return path


def get_file_info(file_record):
    """
    Convert file database object into
    JSON/export-friendly data.
    """
    return {
        "id": file_record.id,

        "original_name": safe_text(
            getattr(
                file_record,
                "original_name",
                ""
            )
        ),

        "stored_name": safe_text(
            getattr(
                file_record,
                "stored_name",
                ""
            )
        ),

        "file_type": safe_text(
            getattr(
                file_record,
                "file_type",
                ""
            )
        ),

        "file_size": getattr(
            file_record,
            "file_size",
            None
        ),

        "uploaded_at": (
            file_record.uploaded_at.isoformat()
            if getattr(
                file_record,
                "uploaded_at",
                None
            )
            else None
        ),
    }


# =========================================================
# BUILD HIRE RECORD
# =========================================================

def build_hire_record(
    hire,
    user=None
):
    """
    Build complete Hire record.

    Fields here correspond to the actual
    Hire form fields plus admin/system information.
    """

    user = user or User.query.filter_by(
        id=hire.user_id
    ).first()

    files = (
        HireFile.query
        .filter_by(
            hire_id=hire.id
        )
        .order_by(
            HireFile.id.asc()
        )
        .all()
    )

    return {
        "record_type": "hire",

        "id": hire.id,

        "user": get_user_info(user),

        # -----------------------------
        # Hire form fields
        # -----------------------------

        "title": safe_text(
            hire.title
        ),

        "type": safe_text(
            hire.type
        ),

        "category": safe_text(
            hire.category
        ),

        "description": safe_text(
            hire.description
        ),

        "scope": safe_text(
            hire.scope
        ),

        "deliverables": safe_text(
            hire.deliverables
        ),

        "budget": (
            float(hire.budget)
            if hire.budget is not None
            else None
        ),

        "processing": safe_text(
            hire.processing
        ),

        "standard_timeline": safe_text(
            hire.standard_timeline
        ),

        "urgent_timeline": safe_text(
            hire.urgent_timeline
        ),

        "human_resources": safe_text(
            hire.human_resources
        ),

        "roles": safe_text(
            hire.roles
        ),

        "skills": safe_text(
            hire.skills
        ),

        "tools": safe_text(
            hire.tools
        ),

        "physical_resources": safe_text(
            hire.physical_resources
        ),

        "additional_resources": safe_text(
            hire.additional_resources
        ),

        "location": safe_text(
            hire.location
        ),

        "full_location": safe_text(
            hire.full_location
        ),

        "communication_method": safe_text(
            hire.communication_method
        ),

        "notification_period": safe_text(
            hire.notification_period
        ),

        # -----------------------------
        # Admin/system fields
        # -----------------------------

        "status": (
            safe_text(hire.status)
            or "Pending"
        ),

        "files": [
            get_file_info(file)
            for file in files
        ],

        "created_at": format_datetime(
            hire.created_at
        ),

        "updated_at": format_datetime(
            hire.updated_at
        ),
    }


# =========================================================
# BUILD DELIVER RECORD
# =========================================================

def build_deliver_record(
    deliver,
    user=None
):
    """
    Build complete Deliver record.

    Includes all actual Deliver form fields,
    team members and uploaded files.
    """

    user = user or User.query.filter_by(
        id=deliver.user_id
    ).first()

    team_members = (
        DeliverTeamMember.query
        .filter_by(
            deliver_id=deliver.id
        )
        .order_by(
            DeliverTeamMember.id.asc()
        )
        .all()
    )

    files = (
        DeliverFile.query
        .filter_by(
            deliver_id=deliver.id
        )
        .order_by(
            DeliverFile.id.asc()
        )
        .all()
    )

    return {
        "record_type": "deliver",

        "id": deliver.id,

        "user": get_user_info(user),

        # -----------------------------
        # Basic Information
        # -----------------------------

        "country": safe_text(
            deliver.country
        ),

        "title": safe_text(
            deliver.title
        ),

        "type": safe_text(
            deliver.type
        ),

        "category": safe_text(
            deliver.category
        ),

        # -----------------------------
        # Delivery Details
        # -----------------------------

        "description": safe_text(
            deliver.description
        ),

        "scope": safe_text(
            deliver.scope
        ),

        "deliverables": safe_text(
            deliver.deliverables
        ),

        # -----------------------------
        # Cost & Processing
        # -----------------------------

        "currency": safe_text(
            deliver.currency
        ),

        "amount": (
            float(deliver.amount)
            if deliver.amount is not None
            else None
        ),

        "included_cost": safe_text(
            deliver.included_cost
        ),

        "excluded_cost": safe_text(
            deliver.excluded_cost
        ),

        "processing": safe_text(
            deliver.processing
        ),

        "standard_timeline": safe_text(
            deliver.standard_timeline
        ),

        "urgent_timeline": safe_text(
            deliver.urgent_timeline
        ),

        # -----------------------------
        # Resources & Tools
        # -----------------------------

        "available_resources": safe_text(
            deliver.available_resources
        ),

        "required_resources": safe_text(
            deliver.required_resources
        ),

        "tools": safe_text(
            deliver.tools
        ),

        "physical_resources": safe_text(
            deliver.physical_resources
        ),

        "additional_resources": safe_text(
            deliver.additional_resources
        ),

        # -----------------------------
        # Location & Portfolio
        # -----------------------------

        "city": safe_text(
            deliver.city
        ),

        "full_location": safe_text(
            deliver.full_location
        ),

        "portfolio": safe_text(
            deliver.portfolio
        ),

        # -----------------------------
        # Notification
        # -----------------------------

        "notification_period": safe_text(
            deliver.notification_period
        ),

        "notification_details": safe_text(
            deliver.notification_details
        ),

        # -----------------------------
        # Communication
        # -----------------------------

        "communication_method": safe_text(
            deliver.communication_method
        ),

        "communication_details": safe_text(
            deliver.communication_details
        ),

        # -----------------------------
        # Team Members
        # -----------------------------

        "team_members": [
            {
                "id": member.id,

                "name": safe_text(
                    member.name
                ),

                "role": safe_text(
                    member.role
                ),

                "skills": safe_text(
                    member.skills
                ),

                "experience": safe_text(
                    member.experience
                ),

                "profile": safe_text(
                    member.profile
                ),

                "similar_experience": safe_text(
                    member.similar_experience
                ),
            }

            for member in team_members
        ],

        # -----------------------------
        # Files
        # -----------------------------

        "files": [
            get_file_info(file)
            for file in files
        ],

        # -----------------------------
        # Admin/system
        # -----------------------------

        "status": (
            safe_text(deliver.status)
            or "Pending"
        ),

        "created_at": format_datetime(
            deliver.created_at
        ),

        "updated_at": format_datetime(
            deliver.updated_at
        ),
    }


# =========================================================
# GET EXPORT RECORDS
# =========================================================

def get_export_records(
    record_type="all",
    search=""
):
    """
    Load records for Excel/CSV export.
    """

    search = safe_text(
        search
    ).strip()

    pattern = (
        f"%{search}%"
        if search
        else None
    )

    hires = []
    delivers = []

    # =====================================================
    # HIRE
    # =====================================================

    if record_type in {
        "all",
        "hire"
    }:

        query = (
            db.session
            .query(
                HireRequirement,
                User
            )
            .join(
                User,
                HireRequirement.user_id
                == User.id
            )
        )

        if pattern:

            query = query.filter(
                or_(
                    HireRequirement.title.ilike(
                        pattern
                    ),

                    HireRequirement.category.ilike(
                        pattern
                    ),

                    HireRequirement.type.ilike(
                        pattern
                    ),

                    HireRequirement.description.ilike(
                        pattern
                    ),

                    HireRequirement.scope.ilike(
                        pattern
                    ),

                    HireRequirement.deliverables.ilike(
                        pattern
                    ),

                    HireRequirement.location.ilike(
                        pattern
                    ),

                    HireRequirement.full_location.ilike(
                        pattern
                    ),

                    HireRequirement.human_resources.ilike(
                        pattern
                    ),

                    HireRequirement.roles.ilike(
                        pattern
                    ),

                    HireRequirement.skills.ilike(
                        pattern
                    ),

                    HireRequirement.tools.ilike(
                        pattern
                    ),

                    HireRequirement.physical_resources.ilike(
                        pattern
                    ),

                    HireRequirement.additional_resources.ilike(
                        pattern
                    ),

                    User.name.ilike(
                        pattern
                    ),

                    User.email.ilike(
                        pattern
                    ),
                )
            )

        hires = [
            build_hire_record(
                hire,
                user
            )

            for hire, user in (
                query
                .order_by(
                    HireRequirement.created_at.desc()
                )
                .all()
            )
        ]

    # =====================================================
    # DELIVER
    # =====================================================

    if record_type in {
        "all",
        "deliver"
    }:

        query = (
            db.session
            .query(
                DeliverProposal,
                User
            )
            .join(
                User,
                DeliverProposal.user_id
                == User.id
            )
        )

        if pattern:

            query = query.filter(
                or_(
                    DeliverProposal.title.ilike(
                        pattern
                    ),

                    DeliverProposal.category.ilike(
                        pattern
                    ),

                    DeliverProposal.type.ilike(
                        pattern
                    ),

                    DeliverProposal.description.ilike(
                        pattern
                    ),

                    DeliverProposal.scope.ilike(
                        pattern
                    ),

                    DeliverProposal.deliverables.ilike(
                        pattern
                    ),

                    DeliverProposal.country.ilike(
                        pattern
                    ),

                    DeliverProposal.city.ilike(
                        pattern
                    ),

                    DeliverProposal.full_location.ilike(
                        pattern
                    ),

                    DeliverProposal.available_resources.ilike(
                        pattern
                    ),

                    DeliverProposal.required_resources.ilike(
                        pattern
                    ),

                    DeliverProposal.tools.ilike(
                        pattern
                    ),

                    DeliverProposal.physical_resources.ilike(
                        pattern
                    ),

                    DeliverProposal.additional_resources.ilike(
                        pattern
                    ),

                    User.name.ilike(
                        pattern
                    ),

                    User.email.ilike(
                        pattern
                    ),
                )
            )

        delivers = [
            build_deliver_record(
                deliver,
                user
            )

            for deliver, user in (
                query
                .order_by(
                    DeliverProposal.created_at.desc()
                )
                .all()
            )
        ]

    return hires, delivers


# =========================================================
# ADMIN RECORDS
# =========================================================

@admin_bp.route(
    "/records",
    methods=["GET"]
)
@admin_required
def get_all_records():

    try:

        search = request.args.get(
            "search",
            "",
            type=str
        ).strip()

        record_type = normalize_record_type(
            request.args.get(
                "type",
                "all",
                type=str
            )
        )

        status_filter = request.args.get(
            "status",
            "",
            type=str
        ).strip()

        page, per_page = parse_page()

        pattern = (
            f"%{search}%"
            if search
            else None
        )

        hires = []
        delivers = []

        # =================================================
        # HIRE RECORDS
        # =================================================

        if record_type in {
            "all",
            "hire"
        }:

            query = (
                db.session
                .query(
                    HireRequirement,
                    User
                )
                .join(
                    User,
                    HireRequirement.user_id
                    == User.id
                )
            )

            if pattern:

                query = query.filter(
                    or_(
                        HireRequirement.title.ilike(
                            pattern
                        ),

                        HireRequirement.category.ilike(
                            pattern
                        ),

                        HireRequirement.type.ilike(
                            pattern
                        ),

                        HireRequirement.description.ilike(
                            pattern
                        ),

                        HireRequirement.location.ilike(
                            pattern
                        ),

                        HireRequirement.full_location.ilike(
                            pattern
                        ),

                        HireRequirement.human_resources.ilike(
                            pattern
                        ),

                        HireRequirement.roles.ilike(
                            pattern
                        ),

                        HireRequirement.skills.ilike(
                            pattern
                        ),

                        HireRequirement.tools.ilike(
                            pattern
                        ),

                        HireRequirement.physical_resources.ilike(
                            pattern
                        ),

                        HireRequirement.additional_resources.ilike(
                            pattern
                        ),

                        User.name.ilike(
                            pattern
                        ),

                        User.email.ilike(
                            pattern
                        ),
                    )
                )

            if status_filter in ALLOWED_STATUSES:

                query = query.filter(
                    HireRequirement.status
                    == status_filter
                )

            hires = [
                build_hire_record(
                    hire,
                    user
                )

                for hire, user in (
                    query
                    .order_by(
                        HireRequirement.created_at.desc()
                    )
                    .all()
                )
            ]

        # =================================================
        # DELIVER RECORDS
        # =================================================

        if record_type in {
            "all",
            "deliver"
        }:

            query = (
                db.session
                .query(
                    DeliverProposal,
                    User
                )
                .join(
                    User,
                    DeliverProposal.user_id
                    == User.id
                )
            )

            if pattern:

                query = query.filter(
                    or_(
                        DeliverProposal.title.ilike(
                            pattern
                        ),

                        DeliverProposal.category.ilike(
                            pattern
                        ),

                        DeliverProposal.type.ilike(
                            pattern
                        ),

                        DeliverProposal.description.ilike(
                            pattern
                        ),

                        DeliverProposal.country.ilike(
                            pattern
                        ),

                        DeliverProposal.city.ilike(
                            pattern
                        ),

                        DeliverProposal.full_location.ilike(
                            pattern
                        ),

                        DeliverProposal.available_resources.ilike(
                            pattern
                        ),

                        DeliverProposal.required_resources.ilike(
                            pattern
                        ),

                        DeliverProposal.tools.ilike(
                            pattern
                        ),

                        DeliverProposal.physical_resources.ilike(
                            pattern
                        ),

                        DeliverProposal.additional_resources.ilike(
                            pattern
                        ),

                        User.name.ilike(
                            pattern
                        ),

                        User.email.ilike(
                            pattern
                        ),
                    )
                )

            if status_filter in ALLOWED_STATUSES:

                query = query.filter(
                    DeliverProposal.status
                    == status_filter
                )

            delivers = [
                build_deliver_record(
                    deliver,
                    user
                )

                for deliver, user in (
                    query
                    .order_by(
                        DeliverProposal.created_at.desc()
                    )
                    .all()
                )
            ]

        # =================================================
        # COMBINE
        # =================================================

        records = sorted(
            hires + delivers,
            key=lambda x: x.get(
                "created_at",
                ""
            ),
            reverse=True
        )

        total = len(records)

        total_pages = max(
            1,
            (
                total
                + per_page
                - 1
            )
            // per_page
        )

        page = min(
            page,
            total_pages
        )

        start = (
            page - 1
        ) * per_page

        paginated = records[
            start:start + per_page
        ]

        # =================================================
        # STATUS BREAKDOWN
        # =================================================

        breakdown = {
            status: 0
            for status in ALLOWED_STATUSES
        }

        for record in records:

            status = record.get(
                "status"
            )

            if status in breakdown:

                breakdown[status] += 1

        return {
            "status": "success",

            "message":
                "Admin records loaded successfully",

            "records": paginated,

            "pagination": {
                "page": page,
                "per_page": per_page,
                "total_records": total,
                "total_pages": total_pages,
                "has_next":
                    page < total_pages,
                "has_previous":
                    page > 1,
            },

            "summary": {
                "total_hires": len(hires),
                "total_delivers": len(delivers),
                "total_records": total,
                "status_breakdown":
                    breakdown,
            },
        }, 200

    except Exception as e:

        db.session.rollback()

        return {
            "status": "error",
            "message":
                "Failed to load admin records",
            "error": str(e),
        }, 500


# =========================================================
# ADMIN OVERVIEW
# =========================================================

@admin_bp.route(
    "/overview",
    methods=["GET"]
)
@admin_required
def admin_overview():

    try:

        total_hires = (
            HireRequirement.query.count()
        )

        total_delivers = (
            DeliverProposal.query.count()
        )

        hire_status = {
            status:
                HireRequirement.query
                .filter_by(
                    status=status
                )
                .count()

            for status in ALLOWED_STATUSES
        }

        deliver_status = {
            status:
                DeliverProposal.query
                .filter_by(
                    status=status
                )
                .count()

            for status in ALLOWED_STATUSES
        }

        return {

            "status": "success",

            "message":
                "Admin overview loaded successfully",

            "overview": {

                "total_users":
                    User.query.count(),

                "total_hires":
                    total_hires,

                "total_delivers":
                    total_delivers,

                "total_records":
                    total_hires
                    + total_delivers,

                "total_files":
                    HireFile.query.count()
                    + DeliverFile.query.count(),

                "hire_status":
                    hire_status,

                "deliver_status":
                    deliver_status,

                "status_breakdown": {
                    status:
                        hire_status[status]
                        + deliver_status[status]

                    for status in ALLOWED_STATUSES
                },
            },

        }, 200

    except Exception as e:

        return {
            "status": "error",
            "message":
                "Failed to load admin overview",
            "error": str(e),
        }, 500


# =========================================================
# ADMIN USERS
# =========================================================

@admin_bp.route(
    "/users",
    methods=["GET"]
)
@admin_required
def get_users():

    try:

        search = request.args.get(
            "search",
            "",
            type=str
        ).strip()

        pattern = (
            f"%{search}%"
            if search
            else None
        )

        query = User.query

        if pattern:

            query = query.filter(
                or_(
                    User.name.ilike(
                        pattern
                    ),

                    User.email.ilike(
                        pattern
                    ),

                    User.phone.ilike(
                        pattern
                    ),
                )
            )

        users = (
            query
            .order_by(
                User.id.desc()
            )
            .all()
        )

        return {

            "status": "success",

            "message":
                "Users loaded successfully",

            "users": [
                get_user_info(user)
                for user in users
            ],

            "total":
                len(users),

        }, 200

    except Exception as e:

        return {
            "status": "error",
            "message":
                "Failed to load users",
            "error": str(e),
        }, 500


# =========================================================
# UPDATE STATUS
# =========================================================

@admin_bp.route(
    "/records/<string:record_type>/<int:record_id>/status",
    methods=["PUT"]
)
@admin_required
def update_record_status(
    record_type,
    record_id
):

    try:

        data = (
            request.get_json(
                silent=True
            )
            or {}
        )

        new_status = safe_text(
            data.get("status")
        ).strip()

        if new_status not in ALLOWED_STATUSES:

            return {
                "status": "error",
                "message":
                    "Invalid status. Allowed statuses: "
                    + ", ".join(
                        ALLOWED_STATUSES
                    ),
            }, 400

        record_type = (
            record_type.lower()
        )

        if record_type == "hire":

            model = HireRequirement

        elif record_type == "deliver":

            model = DeliverProposal

        else:

            return {
                "status": "error",
                "message":
                    "Record type must be hire or deliver",
            }, 400

        record = (
            model.query
            .filter_by(
                id=record_id
            )
            .first()
        )

        if not record:

            return {
                "status": "error",
                "message":
                    "Record not found",
            }, 404

        record.status = new_status

        db.session.commit()

        return {

            "status": "success",

            "message":
                "Record status updated successfully",

            "record_type":
                record_type,

            "record_id":
                record.id,

            "status":
                record.status,

        }, 200

    except Exception as e:

        db.session.rollback()

        return {
            "status": "error",
            "message":
                "Failed to update record status",
            "error": str(e),
        }, 500


# =========================================================
# DELETE RECORD
# =========================================================

@admin_bp.route(
    "/records/<string:record_type>/<int:record_id>",
    methods=["DELETE"]
)
@admin_required
def delete_record(
    record_type,
    record_id
):

    try:

        record_type = (
            record_type.lower()
        )

        physical_files = []

        # =================================================
        # HIRE
        # =================================================

        if record_type == "hire":

            record = (
                HireRequirement.query
                .filter_by(
                    id=record_id
                )
                .first()
            )

            if not record:

                return {
                    "status": "error",
                    "message":
                        "Hire record not found",
                }, 404

            files = (
                HireFile.query
                .filter_by(
                    hire_id=record.id
                )
                .all()
            )

            physical_files = [
                get_file_path(file)
                for file in files
            ]

            (
                HireFile.query
                .filter_by(
                    hire_id=record.id
                )
                .delete(
                    synchronize_session=False
                )
            )

        # =================================================
        # DELIVER
        # =================================================

        elif record_type == "deliver":

            record = (
                DeliverProposal.query
                .filter_by(
                    id=record_id
                )
                .first()
            )

            if not record:

                return {
                    "status": "error",
                    "message":
                        "Deliver record not found",
                }, 404

            files = (
                DeliverFile.query
                .filter_by(
                    deliver_id=record.id
                )
                .all()
            )

            physical_files = [
                get_file_path(file)
                for file in files
            ]

            (
                DeliverTeamMember.query
                .filter_by(
                    deliver_id=record.id
                )
                .delete(
                    synchronize_session=False
                )
            )

            (
                DeliverFile.query
                .filter_by(
                    deliver_id=record.id
                )
                .delete(
                    synchronize_session=False
                )
            )

        else:

            return {
                "status": "error",
                "message":
                    "Record type must be hire or deliver",
            }, 400

        # =================================================
        # DATABASE DELETE
        # =================================================

        db.session.delete(record)

        db.session.commit()

        # =================================================
        # PHYSICAL FILE DELETE
        # =================================================

        for path in physical_files:

            try:

                if (
                    path
                    and os.path.isfile(path)
                ):
                    os.remove(path)

            except Exception:

                pass

        return {

            "status": "success",

            "message":
                "Record deleted successfully",

            "record_type":
                record_type,

            "record_id":
                record_id,

        }, 200

    except Exception as e:

        db.session.rollback()

        return {
            "status": "error",
            "message":
                "Failed to delete record",
            "error": str(e),
        }, 500


# =========================================================
# EXCEL STYLING
# =========================================================

def style_excel_sheet(
    worksheet
):

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="172033"
    )

    header_font = Font(
        bold=True,
        color="FFFFFF"
    )

    # Header styling
    for cell in worksheet[1]:

        cell.fill = header_fill

        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    worksheet.freeze_panes = "A2"

    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    # Body styling
    for row in worksheet.iter_rows():

        for cell in row:

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    # Automatic column widths
    for column_cells in worksheet.columns:

        max_length = 0

        for cell in column_cells:

            value = (
                ""
                if cell.value is None
                else str(cell.value)
            )

            line_lengths = [
                len(line)
                for line in value.split(
                    "\n"
                )
            ]

            if line_lengths:

                max_length = max(
                    max_length,
                    max(line_lengths)
                )

        letter = get_column_letter(
            column_cells[0].column
        )

        worksheet.column_dimensions[
            letter
        ].width = min(
            max(
                max_length + 2,
                12
            ),
            45
        )


# =========================================================
# SECURE EXCEL FILE TOKEN
# =========================================================

def file_link_token(
    record_type,
    file_id
):

    serializer = URLSafeTimedSerializer(
        current_app.config.get(
            "JWT_SECRET_KEY",
            "hire_deliver_file_secret"
        )
    )

    return serializer.dumps(
        {
            "record_type":
                record_type,

            "file_id":
                int(file_id),
        }
    )


def make_file_url(
    record_type,
    file_id
):

    token = file_link_token(
        record_type,
        file_id
    )

    return (
        f"{request.host_url.rstrip('/')}"
        f"/api/admin/files/"
        f"{record_type}/"
        f"{int(file_id)}"
        f"/link?token={token}"
    )


# =========================================================
# ADD FILE ROWS
# =========================================================

def append_file_rows(
    sheet,
    records
):

    for record in records:

        for file in record.get(
            "files",
            []
        ):

            row = (
                sheet.max_row + 1
            )

            sheet.cell(
                row=row,
                column=1,
                value=safe_text(
                    record.get(
                        "record_type",
                        ""
                    )
                ).capitalize()
            )

            sheet.cell(
                row=row,
                column=2,
                value=record.get(
                    "id"
                )
            )

            sheet.cell(
                row=row,
                column=3,
                value=file.get(
                    "id"
                )
            )

            sheet.cell(
                row=row,
                column=4,
                value=file.get(
                    "original_name"
                )
            )

            sheet.cell(
                row=row,
                column=5,
                value=file.get(
                    "file_type"
                )
            )

            sheet.cell(
                row=row,
                column=6,
                value=file.get(
                    "file_size"
                )
            )

            sheet.cell(
                row=row,
                column=7,
                value=file.get(
                    "uploaded_at"
                )
            )

            link_cell = sheet.cell(
                row=row,
                column=8,
                value="Open File"
            )

            link_cell.hyperlink = (
                make_file_url(
                    record.get(
                        "record_type"
                    ),
                    file.get(
                        "id"
                    )
                )
            )

            link_cell.style = "Hyperlink"


# =========================================================
# HIRE EXCEL SHEET
# =========================================================

def add_hire_sheet(
    workbook,
    records
):

    sheet = workbook.create_sheet(
        "Hire Records"
    )

    sheet.append([

        "Record ID",
        "User ID",
        "User Name",
        "User Email",
        "User Phone",

        "Requirement Title",
        "Requirement Type",
        "Category",

        "What Do You Need?",
        "Requirement Scope",
        "Tasks / Deliverables",

        "Total Budget",
        "Processing Type",

        "Standard / Required Timeline",
        "Urgent / Optimal Timeline",

        "Required Human Resources",
        "Required Roles",
        "Required Skills & Expertise",
        "Required Tools / Software",
        "Physical Resources / Equipment",
        "Additional Resources Required",

        "Required City",
        "Full Location / Address",

        "Preferred Communication Method",
        "Notification / Update Period",

        "Uploaded Requirement Files",

        "Status",
        "Created At",
        "Updated At",
    ])

    for record in records:

        user = record.get(
            "user",
            {}
        )

        filenames = "\n".join(
            file.get(
                "original_name",
                ""
            )

            for file in record.get(
                "files",
                []
            )
        )

        sheet.append([

            record.get("id"),

            user.get("id"),

            user.get("name"),

            user.get("email"),

            user.get("phone"),

            record.get("title"),

            record.get("type"),

            record.get("category"),

            record.get("description"),

            record.get("scope"),

            record.get("deliverables"),

            record.get("budget"),

            record.get("processing"),

            record.get(
                "standard_timeline"
            ),

            record.get(
                "urgent_timeline"
            ),

            record.get(
                "human_resources"
            ),

            record.get(
                "roles"
            ),

            record.get(
                "skills"
            ),

            record.get(
                "tools"
            ),

            record.get(
                "physical_resources"
            ),

            record.get(
                "additional_resources"
            ),

            record.get(
                "location"
            ),

            record.get(
                "full_location"
            ),

            record.get(
                "communication_method"
            ),

            record.get(
                "notification_period"
            ),

            filenames,

            record.get(
                "status"
            ),

            record.get(
                "created_at"
            ),

            record.get(
                "updated_at"
            ),
        ])

    style_excel_sheet(
        sheet
    )

    return sheet


# =========================================================
# DELIVER EXCEL SHEET
# =========================================================

def add_deliver_sheet(
    workbook,
    records
):

    sheet = workbook.create_sheet(
        "Deliver Records"
    )

    sheet.append([

        "Record ID",
        "User ID",
        "User Name",
        "User Email",
        "User Phone",

        "Country",

        "Delivery Title",
        "Delivery Type",
        "Delivery Category",

        "Description / What Can Your Team Provide?",
        "Delivery Scope",
        "Deliverables / Task Breakdown",

        "Currency",
        "Amount / Delivery Cost",
        "Included Cost",
        "Excluded Cost",

        "Processing Type",

        "Standard Delivery Timeline",
        "Urgent / Priority Timeline",

        "Available Resources",
        "Required Resources",
        "Available Tools & Software",
        "Physical Resources / Equipment",
        "Additional Resources",

        "Delivery City",
        "Full Delivery Location / Address",

        "Portfolio / Website / Demo Link",

        "Notification / Update Period",
        "Additional Notification Details",

        "Preferred Communication Method",
        "Communication Details",

        "Uploaded Supporting Files",

        "Status",
        "Created At",
        "Updated At",
    ])

    for record in records:

        user = record.get(
            "user",
            {}
        )

        filenames = "\n".join(
            file.get(
                "original_name",
                ""
            )

            for file in record.get(
                "files",
                []
            )
        )

        sheet.append([

            record.get("id"),

            user.get("id"),

            user.get("name"),

            user.get("email"),

            user.get("phone"),

            record.get("country"),

            record.get("title"),

            record.get("type"),

            record.get("category"),

            record.get("description"),

            record.get("scope"),

            record.get("deliverables"),

            record.get("currency"),

            record.get("amount"),

            record.get("included_cost"),

            record.get("excluded_cost"),

            record.get("processing"),

            record.get(
                "standard_timeline"
            ),

            record.get(
                "urgent_timeline"
            ),

            record.get(
                "available_resources"
            ),

            record.get(
                "required_resources"
            ),

            record.get(
                "tools"
            ),

            record.get(
                "physical_resources"
            ),

            record.get(
                "additional_resources"
            ),

            record.get(
                "city"
            ),

            record.get(
                "full_location"
            ),

            record.get(
                "portfolio"
            ),

            record.get(
                "notification_period"
            ),

            record.get(
                "notification_details"
            ),

            record.get(
                "communication_method"
            ),

            record.get(
                "communication_details"
            ),

            filenames,

            record.get(
                "status"
            ),

            record.get(
                "created_at"
            ),

            record.get(
                "updated_at"
            ),
        ])

    style_excel_sheet(
        sheet
    )

    return sheet


# =========================================================
# DELIVER TEAM MEMBERS EXCEL SHEET
# =========================================================

def add_team_sheet(
    workbook,
    records
):

    sheet = workbook.create_sheet(
        "Deliver Team Members"
    )

    sheet.append([

        "Deliver Record ID",
        "Member ID",

        "Member Name",
        "Role",
        "Skills / Expertise",
        "Years of Experience",
        "CV / Profile Documents",
        "Similar Experience",
    ])

    for record in records:

        for member in record.get(
            "team_members",
            []
        ):

            sheet.append([

                record.get(
                    "id"
                ),

                member.get(
                    "id"
                ),

                member.get(
                    "name"
                ),

                member.get(
                    "role"
                ),

                member.get(
                    "skills"
                ),

                member.get(
                    "experience"
                ),

                member.get(
                    "profile"
                ),

                member.get(
                    "similar_experience"
                ),
            ])

    style_excel_sheet(
        sheet
    )

    return sheet


# =========================================================
# UPLOADED FILES EXCEL SHEET
# =========================================================

def add_files_sheet(
    workbook,
    records
):

    sheet = workbook.create_sheet(
        "Uploaded Files"
    )

    sheet.append([

        "Record Type",
        "Record ID",
        "File ID",
        "Original Filename",
        "File Type",
        "File Size",
        "Uploaded At",
        "Open File",
    ])

    append_file_rows(
        sheet,
        records
    )

    style_excel_sheet(
        sheet
    )

    return sheet


# =========================================================
# EXCEL EXPORT
# =========================================================

@admin_bp.route(
    "/records/export",
    methods=["GET"]
)
@admin_required
def export_records():

    try:

        record_type = normalize_record_type(
            request.args.get(
                "type",
                "all",
                type=str
            )
        )

        search = request.args.get(
            "search",
            "",
            type=str
        ).strip()

        # =================================================
        # HIRE EXPORT
        # =================================================

        if record_type == "hire":

            hires, _ = get_export_records(
                "hire",
                search
            )

            workbook = Workbook()

            # Remove default Sheet
            default_sheet = workbook.active
            workbook.remove(
                default_sheet
            )

            add_hire_sheet(
                workbook,
                hires
            )

            add_files_sheet(
                workbook,
                hires
            )

            filename = (
                "Hire_Records.xlsx"
            )

        # =================================================
        # DELIVER EXPORT
        # =================================================

        elif record_type == "deliver":

            _, delivers = get_export_records(
                "deliver",
                search
            )

            workbook = Workbook()

            default_sheet = workbook.active
            workbook.remove(
                default_sheet
            )

            add_deliver_sheet(
                workbook,
                delivers
            )

            add_team_sheet(
                workbook,
                delivers
            )

            add_files_sheet(
                workbook,
                delivers
            )

            filename = (
                "Deliver_Records.xlsx"
            )

        # =================================================
        # ALL EXPORT
        # =================================================

        else:

            hires, delivers = get_export_records(
                "all",
                search
            )

            workbook = Workbook()

            default_sheet = workbook.active
            workbook.remove(
                default_sheet
            )

            add_hire_sheet(
                workbook,
                hires
            )

            add_deliver_sheet(
                workbook,
                delivers
            )

            add_team_sheet(
                workbook,
                delivers
            )

            add_files_sheet(
                workbook,
                hires + delivers
            )

            filename = (
                "HireDeliver_Records.xlsx"
            )

        # =================================================
        # SAVE WORKBOOK
        # =================================================

        output = io.BytesIO()

        workbook.save(
            output
        )

        output.seek(0)

        return send_file(

            output,

            as_attachment=True,

            download_name=filename,

            mimetype=(
                "application/"
                "vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )

    except Exception as e:

        return {
            "status": "error",
            "message":
                "Failed to export records",
            "error": str(e),
        }, 500


# =========================================================
# CSV EXPORT
# =========================================================

@admin_bp.route(
    "/records/export/csv",
    methods=["GET"]
)
@admin_required
def export_records_csv():

    try:

        record_type = normalize_record_type(
            request.args.get(
                "type",
                "all",
                type=str
            )
        )

        search = request.args.get(
            "search",
            "",
            type=str
        ).strip()

        hires, delivers = (
            get_export_records(
                record_type,
                search
            )
        )

        records = sorted(
            hires + delivers,
            key=lambda x:
                x.get(
                    "created_at",
                    ""
                ),
            reverse=True
        )

        output = io.StringIO()

        writer = csv.writer(
            output
        )

        writer.writerow([

            "Record ID",
            "Record Type",

            "User ID",
            "User Name",
            "User Email",
            "User Phone",

            "Title",
            "Type",
            "Category",

            "Description",
            "Scope",
            "Deliverables",

            "Budget / Amount",
            "Currency",

            "Processing",

            "Standard Timeline",
            "Urgent Timeline",

            "Location",
            "Full Location",

            "Status",
            "Created At",
            "Updated At",
        ])

        for record in records:

            user = record.get(
                "user",
                {}
            )

            if (
                record.get(
                    "record_type"
                )
                == "hire"
            ):

                amount = record.get(
                    "budget"
                )

                currency = "PKR"

                location = record.get(
                    "location"
                )

            else:

                amount = record.get(
                    "amount"
                )

                currency = record.get(
                    "currency"
                )

                location = record.get(
                    "city"
                )

            writer.writerow([

                record.get(
                    "id"
                ),

                record.get(
                    "record_type"
                ),

                user.get(
                    "id"
                ),

                user.get(
                    "name"
                ),

                user.get(
                    "email"
                ),

                user.get(
                    "phone"
                ),

                record.get(
                    "title"
                ),

                record.get(
                    "type"
                ),

                record.get(
                    "category"
                ),

                record.get(
                    "description"
                ),

                record.get(
                    "scope"
                ),

                record.get(
                    "deliverables"
                ),

                amount,

                currency,

                record.get(
                    "processing"
                ),

                record.get(
                    "standard_timeline"
                ),

                record.get(
                    "urgent_timeline"
                ),

                location,

                record.get(
                    "full_location"
                ),

                record.get(
                    "status"
                ),

                record.get(
                    "created_at"
                ),

                record.get(
                    "updated_at"
                ),
            ])

        return send_file(

            io.BytesIO(
                output
                .getvalue()
                .encode(
                    "utf-8-sig"
                )
            ),

            as_attachment=True,

            download_name=(
                "HireDeliver_Records.csv"
            ),

            mimetype=(
                "text/csv; charset=utf-8"
            ),
        )

    except Exception as e:

        return {
            "status": "error",
            "message":
                "Failed to export CSV",
            "error": str(e),
        }, 500


# =========================================================
# ADMIN FILE OPENING
# =========================================================

def send_admin_file(
    file_record
):

    if not file_record:

        return {
            "status": "error",
            "message":
                "File not found",
        }, 404

    file_path = get_file_path(
        file_record
    )

    if not file_path:

        return {
            "status": "error",
            "message":
                "Invalid file location",
        }, 403

    if not os.path.isfile(
        file_path
    ):

        return {
            "status": "error",
            "message":
                "Physical file not found",
        }, 404

    mimetype = safe_text(
        getattr(
            file_record,
            "file_type",
            ""
        )
    ) or None

    return send_file(

        file_path,

        as_attachment=False,

        download_name=(
            safe_text(
                file_record.original_name
            )
            or "file"
        ),

        mimetype=mimetype,
    )


# =========================================================
# ADMIN PROTECTED HIRE FILE
# =========================================================

@admin_bp.route(
    "/files/hire/<int:file_id>",
    methods=["GET"]
)
@admin_required
def open_hire_file(
    file_id
):

    file_record = (
        HireFile.query
        .filter_by(
            id=file_id
        )
        .first()
    )

    return send_admin_file(
        file_record
    )


# =========================================================
# ADMIN PROTECTED DELIVER FILE
# =========================================================

@admin_bp.route(
    "/files/deliver/<int:file_id>",
    methods=["GET"]
)
@admin_required
def open_deliver_file(
    file_id
):

    file_record = (
        DeliverFile.query
        .filter_by(
            id=file_id
        )
        .first()
    )

    return send_admin_file(
        file_record
    )


# =========================================================
# EXCEL SIGNED HIRE FILE LINK
# =========================================================

@admin_bp.route(
    "/files/hire/<int:file_id>/link",
    methods=["GET"]
)
def open_hire_file_from_excel(
    file_id
):

    token = request.args.get(
        "token",
        "",
        type=str
    )

    serializer = URLSafeTimedSerializer(
        current_app.config.get(
            "JWT_SECRET_KEY",
            "hire_deliver_file_secret"
        )
    )

    try:

        payload = serializer.loads(
            token,
            max_age=FILE_LINK_MAX_AGE
        )

        if (
            payload.get(
                "record_type"
            )
            != "hire"
        ):

            raise BadSignature(
                "Invalid file token"
            )

        if int(
            payload.get(
                "file_id"
            )
        ) != int(file_id):

            raise BadSignature(
                "Invalid file token"
            )

    except SignatureExpired:

        return {
            "status": "error",
            "message":
                "File link has expired. "
                "Export a new Excel file.",
        }, 401

    except Exception:

        return {
            "status": "error",
            "message":
                "Invalid file link.",
        }, 401

    file_record = (
        HireFile.query
        .filter_by(
            id=file_id
        )
        .first()
    )

    return send_admin_file(
        file_record
    )


# =========================================================
# EXCEL SIGNED DELIVER FILE LINK
# =========================================================

@admin_bp.route(
    "/files/deliver/<int:file_id>/link",
    methods=["GET"]
)
def open_deliver_file_from_excel(
    file_id
):

    token = request.args.get(
        "token",
        "",
        type=str
    )

    serializer = URLSafeTimedSerializer(
        current_app.config.get(
            "JWT_SECRET_KEY",
            "hire_deliver_file_secret"
        )
    )

    try:

        payload = serializer.loads(
            token,
            max_age=FILE_LINK_MAX_AGE
        )

        if (
            payload.get(
                "record_type"
            )
            != "deliver"
        ):

            raise BadSignature(
                "Invalid file token"
            )

        if int(
            payload.get(
                "file_id"
            )
        ) != int(file_id):

            raise BadSignature(
                "Invalid file token"
            )

    except SignatureExpired:

        return {
            "status": "error",
            "message":
                "File link has expired. "
                "Export a new Excel file.",
        }, 401

    except Exception:

        return {
            "status": "error",
            "message":
                "Invalid file link.",
        }, 401

    file_record = (
        DeliverFile.query
        .filter_by(
            id=file_id
        )
        .first()
    )

    return send_admin_file(
        file_record
    )